"""Parse user-supplied OHLCV price files (CSV or Excel).

Equity prices on Borsa Istanbul are licensed data, so the deployed engine
never fetches them itself: users analyse their own exports (for example from
their brokerage) instead. International CSV (``,`` separated, ``.``
decimals), Turkish CSV (``;`` separated, ``,`` decimals, Turkish headers) and
Excel workbooks such as Is Yatirim's historical price download are accepted.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from datetime import date, datetime

import pandas as pd

from investing_engine.providers.base import ProviderError, SourceInfo, validate_price_frame

USER_UPLOAD_SOURCE = SourceInfo(
    name="User-supplied file",
    url="",
    terms="Provided by the user from their own licensed source; processed in memory "
    "for this session and never redistributed.",
    attribution="Price data supplied by the user.",
    deployable=True,
)

_HEADER_ALIASES: dict[str, str] = {
    "date": "Date",
    "tarih": "Date",
    "open": "Open",
    "acilis": "Open",
    "high": "High",
    "max": "High",
    "yuksek": "High",
    "en yuksek": "High",
    "low": "Low",
    "min": "Low",
    "dusuk": "Low",
    "en dusuk": "Low",
    "close": "Close",
    "kapanis": "Close",
    "adj close": "Adj Close",
    "volume": "Volume",
    "hacim": "Volume",
}


_XLSX_MAGIC = b"PK\x03\x04"
# Upper bound on the decompressed size of an Excel upload (decompression-bomb guard).
_MAX_XLSX_EXPANSION = 20


def _fold(text: str) -> str:
    """Normalise a header: lower-case, no diacritics, unit suffixes such as ``(TL)`` removed."""
    text = re.sub(r"\s*\(.*?\)", "", text).strip().lower().replace("ı", "i")
    decomposed = unicodedata.normalize("NFKD", text)
    return "".join(c for c in decomposed if not unicodedata.combining(c))


def parse_ohlcv_csv(raw: bytes, *, symbol: str, max_bytes: int, max_rows: int) -> pd.DataFrame:
    """Validate and parse an uploaded CSV or Excel file into a standard price frame.

    Raises:
        ProviderError: if the file is too large, neither UTF-8 CSV text nor an
            Excel workbook, has no recognisable date/close columns, or
            contains too few rows.
    """
    if len(raw) > max_bytes:
        raise ProviderError(f"File exceeds the {max_bytes // 1024} KiB upload limit")
    if raw.startswith(_XLSX_MAGIC):
        frame = _read_xlsx(raw, max_bytes=max_bytes, max_rows=max_rows)
        decimal, dayfirst = ".", True
    else:
        frame, decimal = _read_csv(raw, max_rows=max_rows)
        dayfirst = decimal == ","

    if len(frame) > max_rows:
        raise ProviderError(f"File exceeds the {max_rows}-row limit")

    renamed = {c: _HEADER_ALIASES.get(_fold(str(c))) for c in frame.columns}
    frame = frame.rename(columns={k: v for k, v in renamed.items() if v})
    if "Close" not in frame.columns and "Adj Close" in frame.columns:
        frame = frame.rename(columns={"Adj Close": "Close"})
    if "Date" not in frame.columns or "Close" not in frame.columns:
        raise ProviderError("File must contain at least 'Date' and 'Close' columns")

    for column in ("Open", "High", "Low", "Close", "Volume"):
        if column in frame.columns:
            frame[column] = _to_number(frame[column], decimal=decimal)

    frame["Date"] = pd.to_datetime(frame["Date"], dayfirst=dayfirst, errors="coerce")
    frame = frame.dropna(subset=["Date"]).set_index("Date")
    return validate_price_frame(frame, symbol=symbol)


def _read_csv(raw: bytes, *, max_rows: int) -> tuple[pd.DataFrame, str]:
    """Read CSV text as strings; returns the frame and its decimal separator."""
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ProviderError("File must be UTF-8 encoded CSV text or an Excel workbook") from exc
    if "\x00" in text:
        raise ProviderError("File appears to be binary, not CSV")

    sample = text[:4096]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
    except csv.Error:
        delimiter = ","
    decimal = "," if delimiter == ";" else "."

    try:
        # Read everything as text; numbers are converted later with the
        # locale-specific decimal separator.
        frame = pd.read_csv(
            io.StringIO(text),
            sep=delimiter,
            nrows=max_rows + 1,
            dtype=str,
            engine="python",
        )
    except (pd.errors.ParserError, ValueError) as exc:
        raise ProviderError("Could not parse the CSV file") from exc
    return frame, decimal


def _read_xlsx(raw: bytes, *, max_bytes: int, max_rows: int) -> pd.DataFrame:
    """Read the first worksheet of an Excel workbook as strings."""
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            expanded = sum(info.file_size for info in archive.infolist())
        if expanded > max_bytes * _MAX_XLSX_EXPANSION:
            raise ProviderError("Excel file expands beyond the allowed size")
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except ProviderError:
        raise
    except Exception as exc:  # openpyxl raises a wide range of errors on malformed files
        raise ProviderError("Could not read the Excel file") from exc

    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True, max_row=max_rows + 2)
        header = next(rows, None)
        if header is None:
            raise ProviderError("Excel file is empty")
        columns = [str(c) if c is not None else "" for c in header]
        records = [[_cell_text(v) for v in row] for row in rows]
    finally:
        workbook.close()
    return pd.DataFrame(records, columns=columns, dtype=str)


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def _to_number(column: pd.Series, *, decimal: str) -> pd.Series:
    values = column.astype(str).str.strip()
    if decimal == ",":
        values = values.str.replace(".", "", regex=False).str.replace(",", ".", regex=False)
    else:
        values = values.str.replace(",", "", regex=False)
    return pd.to_numeric(values, errors="coerce")
